from __future__ import annotations

import csv
import hashlib
import io
import os
import re
from typing import Any

import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from openpyxl import load_workbook
from PIL import Image
from rapidocr_onnxruntime import RapidOCR

APP_ORIGIN = os.getenv("APP_ORIGIN", "https://kwinest-boekhouding.onrender.com")
MAX_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(15 * 1024 * 1024)))
MAX_PDF_PAGES = int(os.getenv("MAX_PDF_PAGES", "50"))
MAX_TABLES_PER_PAGE = 12

app = FastAPI(title="Kwinest Document Processor", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[APP_ORIGIN],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "X-Kwinest-Test-Mode"],
)

_ocr: RapidOCR | None = None


def get_ocr() -> RapidOCR:
    global _ocr
    if _ocr is None:
        _ocr = RapidOCR()
    return _ocr


def clean_text(value: str) -> str:
    return (
        str(value or "")
        .replace("\u00a0", " ")
        .replace("\u2010", "-")
        .replace("\u2011", "-")
        .replace("\u2012", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
    )


def collapse_lines(value: str) -> str:
    lines = []
    for raw in clean_text(value).splitlines():
        line = re.sub(r"[ \t]+", " ", raw).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def safe_cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()[:500]


def doc_type_hint(text: str) -> str:
    low = text.lower()
    if re.search(r"\bcredit(?:factuur|nota)|credit note\b", low):
        return "credit_invoice"
    if re.search(r"\b(factuurnummer|factuurnr|invoice number|invoice no\.?|btw|vat)\b", low):
        return "invoice"
    if re.search(r"\b(kassabon|bonnummer|receipt|totaal betaald)\b", low):
        return "receipt"
    if re.search(r"\b(bankafschrift|account statement|iban|saldo)\b", low):
        return "bank_document"
    return "other"


def text_quality(text: str) -> float:
    text = text or ""
    if not text:
        return 0.0
    printable = sum(ch.isprintable() for ch in text)
    alnum = sum(ch.isalnum() for ch in text)
    replacement = text.count("\ufffd") + text.count("�")
    ratio = printable / max(1, len(text))
    density = alnum / max(1, len(text))
    penalty = min(0.45, replacement / max(1, len(text)) * 8)
    return max(0.0, min(1.0, 0.55 * ratio + 0.45 * density - penalty))


def ocr_image(image: Image.Image) -> tuple[str, float, list[dict[str, Any]]]:
    engine = get_ocr()
    result, _ = engine(image)
    if not result:
        return "", 0.0, []
    words: list[dict[str, Any]] = []
    texts: list[str] = []
    confs: list[float] = []
    for row in result:
        if not row or len(row) < 3:
            continue
        box, text, score = row[0], str(row[1] or "").strip(), float(row[2] or 0)
        if not text:
            continue
        texts.append(text)
        confs.append(score)
        words.append({"text": text, "confidence": round(score, 4), "box": box})
    return collapse_lines("\n".join(texts)), (sum(confs) / len(confs) if confs else 0.0), words


def extract_pdf(data: bytes) -> dict[str, Any]:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise HTTPException(422, f"PDF kon niet worden geopend: {exc}")
    if doc.page_count < 1:
        raise HTTPException(422, "PDF bevat geen pagina's.")
    if doc.page_count > MAX_PDF_PAGES:
        raise HTTPException(413, f"PDF bevat meer dan {MAX_PDF_PAGES} pagina's.")

    pages: list[dict[str, Any]] = []
    full_text: list[str] = []
    ocr_pages: list[int] = []
    warnings: list[str] = []

    try:
        plumber = pdfplumber.open(io.BytesIO(data))
    except Exception:
        plumber = None

    try:
        for idx in range(doc.page_count):
            page = doc.load_page(idx)
            text = collapse_lines(page.get_text("text", sort=True))
            words_raw = page.get_text("words", sort=True)
            words = [
                {
                    "x0": round(float(w[0]), 2),
                    "y0": round(float(w[1]), 2),
                    "x1": round(float(w[2]), 2),
                    "y1": round(float(w[3]), 2),
                    "text": safe_cell(w[4]),
                    "block": int(w[5]),
                    "line": int(w[6]),
                }
                for w in words_raw[:6000]
                if len(w) >= 7 and str(w[4]).strip()
            ]
            quality = text_quality(text)
            ocr_text = ""
            ocr_conf = None
            ocr_words: list[dict[str, Any]] = []
            needs_ocr = len(re.sub(r"\s+", "", text)) < 80 or quality < 0.45

            if needs_ocr:
                matrix = fitz.Matrix(2.2, 2.2)
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
                ocr_text, ocr_conf, ocr_words = ocr_image(image)
                if ocr_text:
                    ocr_pages.append(idx + 1)
                    text = collapse_lines(text + ("\n--- OCR ---\n" if text else "") + ocr_text)
                else:
                    warnings.append(f"Pagina {idx + 1}: OCR leverde geen bruikbare tekst op.")

            tables: list[list[list[str]]] = []
            if plumber is not None and idx < len(plumber.pages):
                try:
                    extracted = plumber.pages[idx].extract_tables(
                        {
                            "vertical_strategy": "lines",
                            "horizontal_strategy": "lines",
                            "intersection_tolerance": 5,
                            "snap_tolerance": 4,
                            "join_tolerance": 4,
                        }
                    ) or []
                    if not extracted:
                        extracted = plumber.pages[idx].extract_tables() or []
                    for table in extracted[:MAX_TABLES_PER_PAGE]:
                        rows = [[safe_cell(cell) for cell in row] for row in table if row]
                        if rows:
                            tables.append(rows[:150])
                except Exception:
                    pass

            full_text.append(text)
            pages.append(
                {
                    "pageNumber": idx + 1,
                    "text": text,
                    "digitalTextChars": len(re.sub(r"\s+", "", collapse_lines(page.get_text("text", sort=True)))),
                    "textQuality": round(quality, 4),
                    "source": "pdf_text+ocr" if ocr_text else "pdf_text",
                    "ocrConfidence": round(float(ocr_conf), 4) if ocr_conf is not None else None,
                    "words": words,
                    "ocrWords": ocr_words[:2500],
                    "tables": tables,
                    "width": round(float(page.rect.width), 2),
                    "height": round(float(page.rect.height), 2),
                }
            )
    finally:
        if plumber is not None:
            plumber.close()
        doc.close()

    combined = "\n\n===== PAGINA =====\n\n".join(full_text)
    return {
        "kind": "pdf",
        "pageCount": len(pages),
        "text": combined,
        "pages": pages,
        "tables": [
            {"page": p["pageNumber"], "table": t}
            for p in pages
            for t in p.get("tables", [])
        ],
        "ocrUsed": bool(ocr_pages),
        "ocrPages": ocr_pages,
        "warnings": warnings,
        "documentTypeHint": doc_type_hint(combined),
    }


def extract_image(data: bytes) -> dict[str, Any]:
    try:
        image = Image.open(io.BytesIO(data)).convert("RGB")
    except Exception as exc:
        raise HTTPException(422, f"Afbeelding kon niet worden gelezen: {exc}")
    text, conf, words = ocr_image(image)
    return {
        "kind": "image",
        "pageCount": 1,
        "text": text,
        "pages": [{"pageNumber": 1, "text": text, "source": "ocr", "ocrConfidence": round(conf, 4), "ocrWords": words[:3000], "tables": []}],
        "tables": [],
        "ocrUsed": True,
        "ocrPages": [1],
        "warnings": [] if text else ["OCR kon geen bruikbare tekst herkennen."],
        "documentTypeHint": doc_type_hint(text),
    }


def extract_docx(data: bytes) -> dict[str, Any]:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(422, f"Word-bestand kon niet worden gelezen: {exc}")
    paragraphs = [collapse_lines(p.text) for p in doc.paragraphs if collapse_lines(p.text)]
    tables = []
    for table in doc.tables[:30]:
        rows = [[safe_cell(cell.text) for cell in row.cells] for row in table.rows[:200]]
        if rows:
            tables.append(rows)
    text = "\n".join(paragraphs)
    if tables:
        text += "\n\n" + "\n".join(" | ".join(row) for table in tables for row in table)
    return {"kind": "docx", "pageCount": None, "text": text, "pages": [], "tables": [{"page": None, "table": t} for t in tables], "ocrUsed": False, "ocrPages": [], "warnings": [], "documentTypeHint": doc_type_hint(text)}


def extract_xlsx(data: bytes) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(422, f"Excel-bestand kon niet worden gelezen: {exc}")
    tables = []
    text_lines = []
    for ws in wb.worksheets[:10]:
        rows = []
        for ridx, row in enumerate(ws.iter_rows(values_only=True)):
            if ridx >= 1000:
                break
            vals = [safe_cell(v) for v in row]
            if any(vals):
                rows.append(vals[:40])
        if rows:
            tables.append({"page": ws.title, "table": rows})
            text_lines.append(f"Werkblad: {ws.title}")
            text_lines.extend(" | ".join(r) for r in rows[:250])
    text = "\n".join(text_lines)
    return {"kind": "xlsx", "pageCount": None, "text": text, "pages": [], "tables": tables, "ocrUsed": False, "ocrPages": [], "warnings": [], "documentTypeHint": doc_type_hint(text)}


def extract_csv_bytes(data: bytes) -> dict[str, Any]:
    decoded = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            decoded = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise HTTPException(422, "CSV-tekst kon niet worden gedecodeerd.")
    sample = decoded[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except Exception:
        dialect = csv.excel
        dialect.delimiter = ";"
    rows = []
    for idx, row in enumerate(csv.reader(io.StringIO(decoded), dialect)):
        if idx >= 2000:
            break
        vals = [safe_cell(v) for v in row]
        if any(vals):
            rows.append(vals[:60])
    text = "\n".join(" | ".join(r) for r in rows[:500])
    return {"kind": "csv", "pageCount": None, "text": text, "pages": [], "tables": [{"page": None, "table": rows}], "ocrUsed": False, "ocrPages": [], "warnings": [], "documentTypeHint": doc_type_hint(text)}


@app.get("/health")
def health() -> dict[str, Any]:
    return {"ok": True, "service": "kwinest-document-processor", "version": "1.0.0"}


@app.post("/extract")
async def extract(file: UploadFile = File(...)) -> JSONResponse:
    filename = (file.filename or "document").strip()[:180]
    data = await file.read(MAX_BYTES + 1)
    await file.close()
    if not data:
        raise HTTPException(400, "Bestand is leeg.")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"Bestand is groter dan {MAX_BYTES // 1024 // 1024} MB.")

    ext = os.path.splitext(filename.lower())[1]
    mime = (file.content_type or "").lower()
    digest = hashlib.sha256(data).hexdigest()

    if ext == ".pdf" or mime == "application/pdf" or data[:5] == b"%PDF-":
        result = extract_pdf(data)
    elif ext in {".png", ".jpg", ".jpeg", ".webp"} or mime.startswith("image/"):
        result = extract_image(data)
    elif ext == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        result = extract_docx(data)
    elif ext == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        result = extract_xlsx(data)
    elif ext == ".csv" or mime in {"text/csv", "application/csv", "text/plain"}:
        result = extract_csv_bytes(data)
    else:
        raise HTTPException(415, "Ondersteund: PDF, JPG, PNG, WEBP, DOCX, XLSX en CSV.")

    result.update({"ok": True, "filename": filename, "mimeType": mime, "size": len(data), "sha256": digest})
    return JSONResponse(result)


def _test_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    try:
        for value in pages:
            page = doc.new_page(width=595, height=842)
            y = 60
            for line in value.splitlines():
                page.insert_text((50, y), line, fontsize=11, fontname="helv")
                y += 18
        return doc.tobytes()
    finally:
        doc.close()


def _test_scan_pdf(value: str) -> bytes:
    try:
        from PIL import ImageFont
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
    except Exception:
        font = None
    image = Image.new("RGB", (1240, 1754), "white")
    from PIL import ImageDraw
    ImageDraw.Draw(image).multiline_text((80, 90), value, fill="black", font=font, spacing=18)
    buf = io.BytesIO()
    image.save(buf, format="PDF", resolution=150.0)
    return buf.getvalue()


@app.get("/self-test")
def self_test(token: str = "") -> dict[str, Any]:
    if not os.getenv("SELFTEST_TOKEN") or token != os.getenv("SELFTEST_TOKEN"):
        raise HTTPException(404, "Not found")

    fixtures = {
        "normal": _test_pdf(["Sligro Food Group B.V.\nFactuurnummer INV-2026-1001\nFactuurdatum 24-09-2026\nSubtotaal 200,00\nBTW 21% 42,00\nTotaal 242,00"]),
        "mixed_vat": _test_pdf(["Food Supplier B.V.\nFactuurnummer MIX-2026-22\nBTW 9% 9,00\nBTW 21% 21,00\nTotaal 230,00"]),
        "multipage": _test_pdf(["MultiPage B.V.\nFactuurnummer MP-2026-77\nPagina 1 van 2", "Pagina 2 van 2\nSubtotaal 750,00\nBTW 157,50\nTotaal 907,50"]),
        "supplier_customer": _test_pdf(["LEVERANCIER\nAlpha Office B.V.\nFACTUUR AAN\nKwinest\nFactuurnummer ALPHA-88\nTotaal 121,00"]),
        "credit": _test_pdf(["Beta Supplies B.V.\nCREDITNOTA\nCreditfactuur CR-2026-09\nTotaal credit 121,00"]),
        "dutch_money": _test_pdf(["Gamma Services B.V.\nFactuurnummer NL-124995\nSubtotaal 1.033,02\nBTW 216,93\nTotaal te betalen 1.249,95"]),
        "missing": _test_pdf(["Onbekende Dienstverlener\nFactuur\nOmschrijving Advieswerkzaamheden\nTotaal 99,00"]),
        "scan": _test_scan_pdf("Scan Supplier B.V.\nKvK 99887766\nFACTUUR\nFactuurnummer SCAN-2026-15\nFactuurdatum 23-09-2026\nSubtotaal 300,00\nBTW 21% 63,00\nTotaal 363,00"),
    }

    results: dict[str, Any] = {}
    for name, blob in fixtures.items():
        result = extract_pdf(blob)
        text = result["text"]
        if name == "normal":
            ok = "INV-2026-1001" in text and "242,00" in text
        elif name == "mixed_vat":
            ok = "9%" in text and "21%" in text and "230,00" in text
        elif name == "multipage":
            ok = result["pageCount"] == 2 and "907,50" in text
        elif name == "supplier_customer":
            ok = "Alpha Office" in text and "Kwinest" in text
        elif name == "credit":
            ok = result["documentTypeHint"] == "credit_invoice" and "CR-2026-09" in text
        elif name == "dutch_money":
            ok = "1.249,95" in text
        elif name == "missing":
            ok = "99,00" in text and "KvK" not in text
        else:
            compact = re.sub(r"[^A-Z0-9]", "", text.upper())
            ok = result["ocrUsed"] and "SCAN" in compact and "363" in compact
        results[name] = {
            "ok": bool(ok),
            "pageCount": result.get("pageCount"),
            "ocrUsed": result.get("ocrUsed"),
            "ocrPages": result.get("ocrPages"),
            "documentTypeHint": result.get("documentTypeHint"),
            "textPreview": text[:240],
        }
    return {"ok": all(v["ok"] for v in results.values()), "cases": results}
